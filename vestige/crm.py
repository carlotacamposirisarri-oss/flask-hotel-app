"""CRM ingestion: CSV / Excel exports from Salesforce, HubSpot, SAP, etc.

Flow: parse upload -> stash rows server-side under a token -> Vestige confirms
the column mapping -> ingest (upsert by email, else by name, within the brand).
"""
import csv
import io
import json
import os
import re
import secrets
from collections import Counter

from flask import current_app

from . import db
from .prompts import INTEREST_CLUSTERS, PHASES
from .segment import parse_date

FIELDS = {
    "": "— ignore —",
    "name": "Full name",
    "first_name": "First name",
    "last_name": "Last name",
    "email": "Email",
    "country": "Country",
    "city": "City",
    "tier": "Tier / segment",
    "archetype": "Archetype",
    "phase": "Relationship phase",
    "engagement_score": "Engagement score (0-100)",
    "lifetime_value": "Lifetime value / total spend",
    "interests": "Interests (comma separated)",
    "purchase_categories": "Product categories bought (comma separated)",
    "purchase_count": "Number of purchases",
    "first_purchase": "First purchase / client since",
    "last_purchase": "Last purchase date",
    "avg_ticket": "Average ticket",
    "social_footprint": "Social footprint",
    "cross_purchases": "Cross-sector purchases",
    "behavioral_signals": "Behavioural signals / last activity",
    "notes": "Notes",
}

SYNONYMS = {
    "name": ["name", "full name", "fullname", "contact name", "client name", "customer name", "account name"],
    "first_name": ["first name", "firstname", "given name", "first"],
    "last_name": ["last name", "lastname", "surname", "family name", "last"],
    "email": ["email", "e-mail", "email address", "mail"],
    "country": ["country", "mailing country", "billing country", "nationality", "country/region"],
    "city": ["city", "mailing city", "billing city", "town"],
    "tier": ["tier", "segment", "client tier", "vip level", "loyalty tier", "status level", "customer segment"],
    "archetype": ["archetype", "persona", "profile type"],
    "phase": ["phase", "stage", "lifecycle stage", "relationship stage", "lead status", "status"],
    "engagement_score": ["engagement", "engagement score", "score", "hubspot score", "lead score"],
    "lifetime_value": ["lifetime value", "ltv", "clv", "total spend", "total revenue", "revenue", "spend",
                       "total purchases", "amount"],
    "interests": ["interests", "hobbies", "passions", "tags", "interest"],
    "purchase_categories": ["product categories", "categories", "categories purchased", "collections", "product lines"],
    "purchase_count": ["purchase count", "number of purchases", "orders", "number of orders", "transactions",
                       "order count", "nb purchases"],
    "first_purchase": ["first purchase", "first purchase date", "client since", "customer since", "first order date"],
    "last_purchase": ["last purchase", "last purchase date", "last order date", "last order", "last transaction"],
    "avg_ticket": ["average ticket", "avg ticket", "average order value", "aov", "average basket"],
    "social_footprint": ["social", "social media", "instagram", "linkedin", "social footprint"],
    "cross_purchases": ["cross purchases", "other brands", "brands owned", "purchases elsewhere"],
    "behavioral_signals": ["last activity", "last visit", "signals", "behaviour", "behavior"],
    "notes": ["notes", "description", "comments", "remarks", "advisor notes"],
}


def _norm(s):
    return re.sub(r"[^a-z0-9/]+", " ", str(s).lower()).strip()


def auto_map(headers):
    mapping, used = {}, set()
    for h in headers:
        n = _norm(h)
        best = ""
        for field, words in SYNONYMS.items():
            if field in used:
                continue
            if n in words:
                best = field
                break
            if not best and any(w in n for w in words if len(w) > 3):
                best = field
        mapping[h] = best
        if best:
            used.add(best)
    return mapping


def parse_upload(storage):
    filename = storage.filename or "upload"
    data = storage.read()
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return [], []
        headers = [str(h or f"Column {i + 1}") for i, h in enumerate(rows[0])]
        body = [{headers[i]: ("" if v is None else str(v)) for i, v in enumerate(r) if i < len(headers)}
                for r in rows[1:] if any(v not in (None, "") for v in r)]
        return headers, body
    text = data.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    headers = reader.fieldnames or []
    return headers, [dict(r) for r in reader if any((v or "").strip() for v in r.values())]


def _stash_dir():
    path = os.path.join(current_app.instance_path, "uploads")
    os.makedirs(path, exist_ok=True)
    return path


def stash(headers, rows, filename, kind="clients"):
    token = secrets.token_hex(12)
    with open(os.path.join(_stash_dir(), token + ".json"), "w", encoding="utf-8") as fh:
        json.dump({"headers": headers, "rows": rows, "filename": filename, "kind": kind}, fh)
    return token


def load_stash(token):
    if not re.fullmatch(r"[0-9a-f]{24}", token or ""):
        return None
    try:
        with open(os.path.join(_stash_dir(), token + ".json"), encoding="utf-8") as fh:
            return json.load(fh)
    except OSError:
        return None


def drop_stash(token):
    try:
        os.remove(os.path.join(_stash_dir(), token + ".json"))
    except OSError:
        pass


def _number(v):
    s = re.sub(r"[^0-9.,\-]", "", str(v or ""))
    if not s:
        return 0.0
    if "," in s and "." in s:
        s = s.replace(",", "") if s.rfind(".") > s.rfind(",") else s.replace(".", "").replace(",", ".")
    elif s.count(",") == 1 and len(s.split(",")[1]) <= 2:
        s = s.replace(",", ".")
    else:
        s = s.replace(",", "")
    try:
        return float(s)
    except ValueError:
        return 0.0


CLUSTER_LOOKUP = {c.lower(): c for c in INTEREST_CLUSTERS}
CLUSTER_LOOKUP.update({"art": "Arts", "wine": "Wine & Spirits", "spirits": "Wine & Spirits", "yachting": "Sailing",
                       "golf": "Sport", "tennis": "Sport", "polo": "Equestrian", "horses": "Equestrian",
                       "cars": "Motorsport", "f1": "Motorsport", "food": "Gastronomy", "dining": "Gastronomy",
                       "watches": "Collectibles", "architecture": "Design & Architecture", "design": "Design & Architecture",
                       "charity": "Philanthropy", "property": "Real Estate", "spa": "Wellness", "opera": "Music"})


def _interests(v):
    out = []
    for part in re.split(r"[,;|/]", str(v or "")):
        p = part.strip()
        if p:
            label = CLUSTER_LOOKUP.get(p.lower(), p.title())
            if label not in out:
                out.append(label)
    return out


def _phase(v):
    v = str(v or "").strip()
    for p in PHASES:
        if p.lower() == v.lower():
            return p
    low = v.lower()
    if any(w in low for w in ("risk", "churn", "lapsed")):
        return "At risk"
    if any(w in low for w in ("inactive", "dormant")):
        return "Dormant"
    if any(w in low for w in ("vip", "loyal", "customer")):
        return "Loyal"
    if any(w in low for w in ("lead", "prospect", "new")):
        return "Discovery"
    return v or "Discovery"


def ingest(brand_id, rows, mapping):
    imported = 0
    for r in rows:
        rec = {}
        for col, field in mapping.items():
            if field and r.get(col) not in (None, ""):
                rec[field] = str(r[col]).strip()
        name = rec.pop("name", "") or " ".join(x for x in (rec.pop("first_name", ""), rec.pop("last_name", "")) if x)
        rec.pop("first_name", None)
        rec.pop("last_name", None)
        if not name:
            continue
        values = {k: v for k, v in rec.items() if k in FIELDS}
        if "engagement_score" in values:
            values["engagement_score"] = max(0, min(100, int(_number(values["engagement_score"]))))
        if "lifetime_value" in values:
            values["lifetime_value"] = _number(values["lifetime_value"])
        if "interests" in values:
            values["interests"] = _interests(values["interests"])
        if "purchase_categories" in values:
            values["purchase_categories"] = [x.strip().title() for x in re.split(r"[,;|]", values["purchase_categories"]) if x.strip()]
        if "purchase_count" in values:
            values["purchase_count"] = int(_number(values["purchase_count"]))
        if "avg_ticket" in values:
            values["avg_ticket"] = _number(values["avg_ticket"])
        for f in ("first_purchase", "last_purchase"):
            if f in values:
                d = parse_date(values[f])
                values[f] = d.isoformat() if d else ""
        if "phase" in values:
            values["phase"] = _phase(values["phase"])
        email = values.get("email", "")
        existing = (db.query("SELECT id FROM clients WHERE brand_id = ? AND email = ? AND email != ''",
                             (brand_id, email), one=True) if email else None) or \
            db.query("SELECT id FROM clients WHERE brand_id = ? AND name = ?", (brand_id, name), one=True)
        values.update(name=name, updated_at=db.now())
        if existing:
            db.update("clients", existing["id"], values)
        else:
            values.update(brand_id=brand_id, source="crm", created_at=db.now())
            db.insert("clients", values)
        imported += 1
    return imported


# ------------------------------------------------------- transaction files ----
# One row per purchase. Aggregated into each client's purchase history.

TX_FIELDS = {
    "": "— ignore —",
    "name": "Client full name",
    "first_name": "Client first name",
    "last_name": "Client last name",
    "email": "Client email",
    "client_ref": "Client ID / reference",
    "date": "Purchase date",
    "amount": "Amount",
    "category": "Product category / collection",
    "country": "Country / store country",
}
TX_SYNONYMS = {
    "client_ref": ["client id", "customer id", "contact id", "account id", "client ref", "customer number"],
    "name": ["name", "full name", "client name", "customer name", "contact name"],
    "first_name": ["first name", "firstname"],
    "last_name": ["last name", "lastname", "surname"],
    "email": ["email", "e-mail", "email address"],
    "date": ["date", "purchase date", "order date", "transaction date", "invoice date"],
    "amount": ["amount", "total", "net amount", "value", "price", "revenue", "sales"],
    "category": ["category", "product category", "collection", "product line", "family", "department"],
    "country": ["country", "store country", "boutique country"],
}


def auto_map_tx(headers):
    mapping, used = {}, set()
    for h in headers:
        n = _norm(h)
        best = next((f for f, words in TX_SYNONYMS.items() if f not in used and n in words), "")
        if not best:
            best = next((f for f, words in TX_SYNONYMS.items()
                         if f not in used and any(w in n for w in words if len(w) > 3)), "")
        mapping[h] = best
        if best:
            used.add(best)
    return mapping


def ingest_transactions(brand_id, rows, mapping):
    """Aggregate purchases per client, then upsert purchase history. Returns (clients, rows used)."""
    agg = {}
    used = 0
    for r in rows:
        rec = {f: str(r.get(col, "")).strip() for col, f in mapping.items() if f and r.get(col) not in (None, "")}
        name = rec.get("name") or " ".join(x for x in (rec.get("first_name", ""), rec.get("last_name", "")) if x)
        key = (rec.get("email") or rec.get("client_ref") or name).lower()
        if not key:
            continue
        a = agg.setdefault(key, {"name": name or rec.get("client_ref", "") or rec.get("email", ""), "email": rec.get("email", ""),
                                 "country": rec.get("country", ""), "dates": [], "amounts": [], "cats": Counter()})
        d = parse_date(rec.get("date"))
        if d:
            a["dates"].append(d)
        a["amounts"].append(_number(rec.get("amount")))
        if rec.get("category"):
            a["cats"][rec["category"].strip().title()] += 1
        used += 1
    for a in agg.values():
        values = {
            "purchase_count": len(a["amounts"]),
            "lifetime_value": round(sum(a["amounts"]), 2),
            "avg_ticket": round(sum(a["amounts"]) / len(a["amounts"]), 2) if a["amounts"] else 0,
            "purchase_categories": [c for c, _ in a["cats"].most_common()],
            "first_purchase": min(a["dates"]).isoformat() if a["dates"] else "",
            "last_purchase": max(a["dates"]).isoformat() if a["dates"] else "",
            "updated_at": db.now(),
        }
        existing = (db.query("SELECT id FROM clients WHERE brand_id = ? AND email = ? AND email != ''",
                             (brand_id, a["email"]), one=True) if a["email"] else None) or \
            db.query("SELECT id FROM clients WHERE brand_id = ? AND name = ?", (brand_id, a["name"]), one=True)
        if existing:
            db.update("clients", existing["id"], values)
        elif a["name"]:
            db.insert("clients", {**values, "brand_id": brand_id, "name": a["name"], "email": a["email"],
                                  "country": a["country"], "source": "crm-transactions", "created_at": db.now()})
    return len(agg), used
