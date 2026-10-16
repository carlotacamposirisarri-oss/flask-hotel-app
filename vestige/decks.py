"""Deliverable Studio: assemble decks from work items, AI-compose them, export PPTX."""
import io
from datetime import date

from . import ai, analytics, db
from .prompts import DECK_PROMPT, KINDS, VESTIGE_SYSTEM

SLIDE_TYPES = ["cover", "section", "text", "metrics", "bars", "quote", "closing"]


def _bullets_from_payload(kind, payload, limit=6):
    out = []
    for key, val in payload.items():
        if key in ("month",):
            continue
        label = key.replace("_", " ").capitalize()
        if isinstance(val, list):
            for v in val[:3]:
                if isinstance(v, dict):
                    v = " · ".join(str(x) for x in v.values())
                out.append(f"{label}: {v}")
        elif isinstance(val, dict):
            out.append(f"{label}: " + "; ".join(f"{k} {v}" for k, v in val.items()))
        elif val not in (None, ""):
            out.append(f"{label}: {val}")
    return out[:limit]


def data_slides(brand):
    if not brand:
        x = analytics.cross_brand()
        return [{"type": "bars", "title": "Interest clusters across the Vestige portfolio",
                 "bars": [{"label": k, "value": v} for k, v in x["interests"][:8]]}]
    s = analytics.portfolio_stats(brand["id"])
    return [
        {"type": "metrics", "title": "The portfolio at a glance", "metrics": [
            {"label": "Clients tracked", "value": str(s["count"])},
            {"label": "Average engagement", "value": str(s["avg_engagement"])},
            {"label": "Loyal & ambassadors", "value": str(s["loyal"])},
            {"label": "At risk or dormant", "value": str(s["at_risk"])},
        ]},
        {"type": "bars", "title": "What moves these clients — interest clusters",
         "bars": [{"label": k, "value": v} for k, v in s["interests"][:8]]},
    ]


def assemble(title, brand, items, audience):
    """Deterministic deck: no AI needed, Vestige refines afterwards."""
    subtitle = f"Prepared by Vestige for {brand['name']}" if brand else "Vestige · Internal strategy"
    slides = [{"type": "cover", "title": title, "subtitle": f"{subtitle} · {date.today():%B %Y}"}]
    slides.append({"type": "text", "title": "Executive summary",
                   "body": "The priorities that emerged from our work, in order of impact.",
                   "bullets": [i["title"] for i in sorted(items, key=lambda i: -(i["impact"] * i["urgency"]))[:6]]})
    slides += data_slides(brand)
    by_kind = {}
    for it in items:
        by_kind.setdefault(it["kind"], []).append(it)
    for kind, its in by_kind.items():
        slides.append({"type": "section", "title": KINDS.get(kind, {}).get("plural", kind.title()),
                       "subtitle": f"{len(its)} recommendation{'s' if len(its) != 1 else ''}"})
        for it in its:
            body = it["summary"] + (f"\n\nVestige note: {it['annotation']}" if it["annotation"] else "")
            slides.append({"type": "text", "title": it["title"], "body": body,
                           "bullets": _bullets_from_payload(kind, it["payload"])})
    slides.append({"type": "closing", "title": "Next steps", "bullets": [
        "Validate priorities together", "Assign owners and timing", "Activate the first initiative",
        "Review signals and impact in 90 days"]})
    return slides


def compose_with_ai(title, brand, items, audience, provider=None):
    context = analytics.brand_codes(brand) if brand else "## Audience: Vestige internal leadership"
    work = "\n\n".join(
        f"[{KINDS.get(i['kind'], {}).get('label', i['kind'])}] {i['title']}\n{i['summary']}\n"
        f"Details: {i['payload']}\nVestige annotation: {i['annotation']}" for i in items)
    data, model = ai.complete_json(VESTIGE_SYSTEM, DECK_PROMPT.format(
        context=context, title=title, audience=audience, work=work), provider)
    slides = [s for s in data.get("slides", []) if isinstance(s, dict) and s.get("type") in SLIDE_TYPES]
    if not slides:
        raise ai.AIError("The model returned no usable slides.")
    # Data slides are computed, not generated: insert real numbers after the summary.
    slides[2:2] = data_slides(brand)
    return slides, model


# ------------------------------------------------------------------ PPTX ----

def _rgb(hex_color):
    from pptx.dml.color import RGBColor

    h = hex_color.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def to_pptx(deck, accent="#8C6A4A"):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches, Pt

    bone, charcoal, sand = _rgb("#F6F2EA"), _rgb("#2B2724"), _rgb("#E6DCCB")
    acc = _rgb(accent or "#8C6A4A")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    def bg(slide, color):
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def text(slide, x, y, w, h, content, size, color, serif=False, bold=False):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        lines = content if isinstance(content, list) else [content]
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = str(line)
            p.font.size = Pt(size)
            p.font.color.rgb = color
            p.font.bold = bold
            p.font.name = "Georgia" if serif else "Helvetica"
            p.space_after = Pt(size * 0.6)
        return tb

    for s in deck["slides"]:
        slide = prs.slides.add_slide(blank)
        t = s.get("type")
        if t in ("cover", "closing", "section"):
            bg(slide, charcoal if t != "section" else acc)
            text(slide, 0.9, 0.6, 6, 0.4, "VESTIGE", 11, sand, bold=True)
            text(slide, 0.9, 2.4, 11.5, 1.8, s.get("title", ""), 44 if t == "cover" else 36, bone, serif=True)
            if s.get("subtitle"):
                text(slide, 0.9, 4.3, 11.5, 0.8, s["subtitle"], 16, sand)
            if s.get("bullets"):
                text(slide, 0.9, 4.0, 11.5, 3, [f"—  {b}" for b in s["bullets"]], 16, sand)
            continue
        bg(slide, bone)
        text(slide, 0.9, 0.5, 6, 0.3, "VESTIGE", 9, acc, bold=True)
        text(slide, 0.9, 0.9, 11.5, 1.2, s.get("title", ""), 30, charcoal, serif=True)
        if t == "metrics":
            ms = s.get("metrics", [])[:4]
            for i, m in enumerate(ms):
                x = 0.9 + i * (11.5 / max(len(ms), 1))
                text(slide, x, 2.8, 2.8, 1.2, m.get("value", ""), 40, acc, serif=True)
                text(slide, x, 4.0, 2.8, 0.8, m.get("label", ""), 13, charcoal)
        elif t == "bars" and s.get("bars"):
            cd = CategoryChartData()
            cd.categories = [b["label"] for b in s["bars"]]
            cd.add_series("Clients", [b["value"] for b in s["bars"]])
            chart = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.9), Inches(2.1),
                                           Inches(11.5), Inches(4.9), cd).chart
            chart.has_legend = False
            plot = chart.plots[0]
            plot.gap_width = 60
            plot.series[0].format.fill.solid()
            plot.series[0].format.fill.fore_color.rgb = acc
            chart.category_axis.reverse_order = True
            chart.category_axis.tick_labels.font.size = Pt(12)
            chart.value_axis.visible = False
            chart.value_axis.has_major_gridlines = False
        elif t == "quote":
            text(slide, 1.5, 2.5, 10.3, 3.5, f"“{s.get('body', '')}”", 28, charcoal, serif=True)
        else:
            if s.get("body"):
                text(slide, 0.9, 2.1, 5.6 if s.get("bullets") else 11.5, 4.8, s["body"].split("\n\n"), 15, charcoal)
            if s.get("bullets"):
                text(slide, 7.0 if s.get("body") else 0.9, 2.1, 5.4 if s.get("body") else 11.5, 4.8,
                     [f"—  {b}" for b in s["bullets"]], 13, charcoal)
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf


def items_for_deck(brand_id=None, scope=None, ids=None):
    if ids:
        marks = ",".join("?" for _ in ids)
        return db.query(f"SELECT * FROM items WHERE id IN ({marks}) ORDER BY kind, priority DESC, id", ids)
    if brand_id:
        return db.query("SELECT * FROM items WHERE brand_id = ? AND status != 'dismissed' ORDER BY kind, id",
                        (brand_id,))
    return db.query("SELECT * FROM items WHERE scope = ? AND status != 'dismissed' ORDER BY kind, id", (scope,))
